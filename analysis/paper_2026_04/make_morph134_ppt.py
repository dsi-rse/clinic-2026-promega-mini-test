#!/usr/bin/env python3
"""Brief PPT: morphology LGBM vs LogReg on IDOR-134 set.

Slides:
  1. Title
  2. morph_lgbm vs morph_logreg — n=134 (line + SD band)
  3. morph_logreg: n=140 vs n=134 — impact of removing stitched organoids
  4. morph_lgbm: n=140 vs n=134 — LGBM barely affected

Usage:
    python3 -m analysis.paper_2026_04.make_morph134_ppt
"""

import io
import json

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from pipeline.data_loader import ANALYSIS_OUTPUT_DIR, DAY_ORDER

OLD_PATH = ANALYSIS_OUTPUT_DIR / "images" / "met_morph_lgbm_logreg_kfold.json"
NEW_PATH = ANALYSIS_OUTPUT_DIR / "images" / "met_morph_lgbm_logreg_kfold_134.json"
OUT_PPT  = "figures/morph_lgbm_vs_logreg_134.pptx"

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
M       = Inches(0.3)

DAYS = list(DAY_ORDER)


# ── helpers ────────────────────────────────────────────────────────────────

def _hex(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

def _new_prs():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def _blank(prs):
    return prs.slide_layouts[6]

def _tb(slide, text, left, top, width, height,
        fontsize=18, bold=False, color="#000000", align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf  = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(fontsize)
    run.font.bold = bold
    run.font.color.rgb = _hex(color)

def _title_slide(prs, title, subtitle="", body=""):
    slide = prs.slides.add_slide(_blank(prs))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _hex("#1F3864")
    _tb(slide, title, M, Inches(1.6), Inches(12.5), Inches(2.8),
        fontsize=38, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
    if subtitle:
        _tb(slide, subtitle, M, Inches(4.2), Inches(12.5), Inches(0.7),
            fontsize=18, color="#BDD7EE", align=PP_ALIGN.CENTER)
    if body:
        _tb(slide, body, M, Inches(5.0), Inches(12.5), Inches(2.0),
            fontsize=14, color="#DDEBF7", align=PP_ALIGN.CENTER)

def _content_slide(prs, heading, fig):
    slide = prs.slides.add_slide(_blank(prs))
    _tb(slide, heading, M, M, Inches(12.5), Inches(0.55),
        fontsize=22, bold=True, color="#1F3864")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    buf.seek(0)
    slide.shapes.add_picture(buf, M, Inches(0.75), width=Inches(12.5))
    plt.close(fig)

def _get(data, day, key):
    r = data.get(day, {}).get(key, {})
    return r.get("balanced_accuracy_mean"), r.get("balanced_accuracy_std")

def _series(data, key):
    xs, ys, lo, hi = [], [], [], []
    for i, d in enumerate(DAYS):
        mn, std = _get(data, d, key)
        if mn is not None:
            xs.append(i); ys.append(mn)
            lo.append(mn - (std or 0)); hi.append(mn + (std or 0))
    return xs, ys, lo, hi

def _style_ax(ax):
    ax.set_xticks(range(len(DAYS)))
    ax.set_xticklabels(DAYS, rotation=45, fontsize=9)
    ax.set_ylim(0.4, 1.05)
    ax.axhline(0.5, color="#cccccc", lw=0.8, ls="--")
    ax.set_ylabel("Balanced Accuracy (mean ± 1 SD)", fontsize=10)
    ax.set_xlabel("Day", fontsize=10)
    ax.grid(axis="y", alpha=0.3)
    ax.spines[["top", "right"]].set_visible(False)


# ── plots ───────────────────────────────────────────────────────────────────

def plot_n134_comparison(new):
    """morph_lgbm vs morph_logreg on n=134."""
    specs = [
        ("morph_lgbm",   "Morph LGBM",   "#1f77b4", "D"),
        ("morph_logreg", "Morph LogReg",  "#d62728", "P"),
    ]
    fig, ax = plt.subplots(figsize=(13, 5))
    for key, label, color, marker in specs:
        xs, ys, lo, hi = _series(new, key)
        if xs:
            ax.plot(xs, ys, marker=marker, color=color, lw=2.5, ms=7, label=label)
            ax.fill_between(xs, lo, hi, color=color, alpha=0.12, lw=0)
    _style_ax(ax)
    ax.set_title("Morphology: LGBM vs LogReg  (IDOR-134, 10×4-fold CV)",
                 fontsize=12, fontweight="bold")
    ax.legend(fontsize=10, loc="upper left")
    plt.tight_layout()
    return fig


def plot_logreg_n140_vs_n134(old, new):
    """morph_logreg: before vs after removing stitched organoids."""
    fig, ax = plt.subplots(figsize=(13, 5))
    for data, label, color, ls, marker in [
        (old, "morph_logreg  n=140 (with stitched)", "#aaaaaa", "--", "o"),
        (new, "morph_logreg  n=134 (stitched removed)", "#d62728", "-",  "P"),
    ]:
        xs, ys, lo, hi = _series(data, "morph_logreg")
        if xs:
            ax.plot(xs, ys, marker=marker, color=color, lw=2.5, ls=ls, ms=7, label=label)
            ax.fill_between(xs, lo, hi, color=color, alpha=0.10, lw=0)
    _style_ax(ax)
    ax.set_title("morph LogReg: effect of removing 6 stitched organoids  (10×4-fold CV)",
                 fontsize=12, fontweight="bold")
    ax.legend(fontsize=10, loc="upper left")
    plt.tight_layout()
    return fig


def plot_lgbm_n140_vs_n134(old, new):
    """morph_lgbm: before vs after — should be nearly unchanged."""
    fig, ax = plt.subplots(figsize=(13, 5))
    for data, label, color, ls, marker in [
        (old, "morph_lgbm  n=140 (with stitched)", "#aaaaaa", "--", "o"),
        (new, "morph_lgbm  n=134 (stitched removed)", "#1f77b4", "-",  "D"),
    ]:
        xs, ys, lo, hi = _series(data, "morph_lgbm")
        if xs:
            ax.plot(xs, ys, marker=marker, color=color, lw=2.5, ls=ls, ms=7, label=label)
            ax.fill_between(xs, lo, hi, color=color, alpha=0.10, lw=0)
    _style_ax(ax)
    ax.set_title("morph LGBM: effect of removing 6 stitched organoids  (10×4-fold CV)",
                 fontsize=12, fontweight="bold")
    ax.legend(fontsize=10, loc="upper left")
    plt.tight_layout()
    return fig


# ── main ────────────────────────────────────────────────────────────────────

def main():
    old = json.loads(OLD_PATH.read_text())
    new = json.loads(NEW_PATH.read_text())

    prs = _new_prs()

    _title_slide(prs,
        title="Morphology Classifier: LGBM vs LogReg",
        subtitle="IDOR-134 set (6 stitched organoids removed)  ·  10×4-fold CV",
        body="Removing stitched organoids reverses the late-day morph LogReg deficit")

    _content_slide(prs,
        "Morph LGBM vs LogReg — IDOR-134 (n=134)",
        plot_n134_comparison(new))

    _content_slide(prs,
        "morph LogReg: n=140 (with stitched) vs n=134 (stitched removed)",
        plot_logreg_n140_vs_n134(old, new))

    _content_slide(prs,
        "morph LGBM: n=140 vs n=134 — minimal impact",
        plot_lgbm_n140_vs_n134(old, new))

    import pathlib
    pathlib.Path("figures").mkdir(exist_ok=True)
    prs.save(OUT_PPT)
    print(f"Saved → {OUT_PPT}")


if __name__ == "__main__":
    main()
